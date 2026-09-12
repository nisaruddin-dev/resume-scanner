"""
Resume Scanner — file extraction and rule-based scoring.

To add a new scoring rule: write `score_<name>(text: str) -> dict` returning
{"score": int, "max": int, "notes": str}, then register it in SCORERS.
Do not modify any other file.
"""
import io
import re
from typing import Callable

import pdfplumber
from docx import Document


# ---------------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------------

def extract_text(file_name: str, file_bytes: bytes) -> dict:
    """Dispatch by extension. Returns {ok, error, raw, summary}."""
    name = file_name.lower()
    try:
        if name.endswith(".pdf"):
            return _extract_pdf(file_bytes)
        if name.endswith(".docx"):
            return _extract_docx(file_bytes)
        return {"ok": False, "error": f"Unsupported file type: {file_name}",
                "raw": {}, "summary": {}}
    except Exception as e:
        return {"ok": False, "error": str(e), "raw": {}, "summary": {}}


def _extract_pdf(file_bytes: bytes) -> dict:
    pages_text: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            pages_text.append(page.extract_text() or "")
    text = "\n".join(pages_text).strip()
    if not text:
        return {"ok": False,
                "error": "PDF has no extractable text. It may be a scanned image — "
                         "try a text-based PDF or a DOCX.",
                "raw": {}, "summary": {}}
    return {"ok": True, "error": None,
            "raw": {"num_pages": len(pages_text)},
            "summary": {"text": text, "num_pages": len(pages_text),
                        "char_count": len(text)}}


def _extract_docx(file_bytes: bytes) -> dict:
    doc = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    text = "\n".join(parts).strip()
    if not text:
        return {"ok": False, "error": "DOCX has no extractable text.",
                "raw": {}, "summary": {}}
    return {"ok": True, "error": None,
            "raw": {"num_paragraphs": len(doc.paragraphs)},
            "summary": {"text": text, "num_paragraphs": len(doc.paragraphs),
                        "char_count": len(text)}}


# ---------------------------------------------------------------------------
# Scoring heuristics
# ---------------------------------------------------------------------------

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"(\+?\d[\d\s\-().]{7,}\d)")
LINKEDIN_RE = re.compile(r"linkedin\.com/in/[\w\-]+", re.IGNORECASE)

SECTION_KEYWORDS = {
    "experience": ["experience", "employment", "work history"],
    "education":  ["education", "academic", "university"],
    "skills":     ["skills", "technologies", "competencies"],
    "projects":   ["projects", "portfolio"],
}

STRONG_VERBS = {
    "led", "built", "designed", "developed", "implemented", "launched", "created",
    "managed", "improved", "reduced", "increased", "delivered", "drove", "achieved",
    "architected", "automated", "optimized", "scaled", "shipped", "owned",
    "spearheaded", "established", "transformed", "accelerated", "generated",
}

QUANT_PATTERNS = [
    re.compile(r"\b\d+(?:\.\d+)?%"),
    re.compile(r"\$\s?\d"),
    re.compile(r"\b\d{2,}\b"),
    re.compile(r"\b\d+(?:k|m|x|bn)\b", re.IGNORECASE),
]


def _bullets(text: str) -> list[str]:
    out = []
    for raw in text.split("\n"):
        line = re.sub(r"^[\-\*\u2022\u25CF\u25AA\u25E6o]\s*", "", raw.strip())
        if len(line) >= 15:
            out.append(line)
    return out


def score_contact(text: str) -> dict:
    has_email = bool(EMAIL_RE.search(text))
    has_phone = bool(PHONE_RE.search(text))
    has_linkedin = bool(LINKEDIN_RE.search(text))
    earned = sum([has_email, has_phone, has_linkedin]) * 3
    missing = []
    if not has_email: missing.append("email")
    if not has_phone: missing.append("phone")
    if not has_linkedin: missing.append("LinkedIn")
    return {"score": earned, "max": 9,
            "notes": "missing: " + ", ".join(missing) if missing else "complete"}


def score_sections(text: str) -> dict:
    low = text.lower()
    found = [k for k, keys in SECTION_KEYWORDS.items()
             if any(kw in low for kw in keys)]
    missing = [k for k in SECTION_KEYWORDS if k not in found]
    return {"score": len(found) * 3, "max": 12,
            "notes": f"found: {', '.join(found) or 'none'}; "
                     f"missing: {', '.join(missing) or 'none'}"}


def score_quantified(text: str) -> dict:
    b = _bullets(text)
    if not b:
        return {"score": 0, "max": 15, "notes": "no bullet points detected"}
    n = sum(1 for line in b if any(p.search(line) for p in QUANT_PATTERNS))
    return {"score": round((n / len(b)) * 15), "max": 15,
            "notes": f"{n}/{len(b)} bullets contain numbers or metrics"}


def score_action_verbs(text: str) -> dict:
    b = _bullets(text)
    if not b:
        return {"score": 0, "max": 15, "notes": "no bullet points detected"}
    n = sum(1 for line in b if line.split()[0].lower().strip(",.:;") in STRONG_VERBS)
    return {"score": round((n / len(b)) * 15), "max": 15,
            "notes": f"{n}/{len(b)} bullets start with a strong action verb"}


def score_length(text: str) -> dict:
    w = len(text.split())
    if 400 <= w <= 900:
        return {"score": 10, "max": 10, "notes": f"{w} words — good length"}
    if 300 <= w < 400 or 900 < w <= 1100:
        return {"score": 6, "max": 10, "notes": f"{w} words — slightly off ideal (aim for 400–900)"}
    return {"score": 3, "max": 10,
            "notes": f"{w} words — {'too short' if w < 300 else 'too long'}"}


def score_formatting(text: str) -> dict:
    flags = []
    if re.search(r"\t{2,}", text) or re.search(r" {4,}\S", text):
        flags.append("possible multi-column layout")
    short_lines = sum(1 for ln in text.split("\n") if 0 < len(ln.strip()) < 8)
    if short_lines > 20:
        flags.append("many very short lines (may be a table)")
    if not flags:
        return {"score": 10, "max": 10, "notes": "clean formatting"}
    return {"score": 4, "max": 10, "notes": "; ".join(flags)}


def score_keywords(text: str, jd: str) -> dict:
    if not jd.strip():
        return {"score": 0, "max": 0, "notes": "no job description provided"}
    stop = {"the", "and", "for", "with", "you", "will", "are", "our", "this",
            "that", "from", "have", "has", "not", "but", "all", "any", "can",
            "able", "into", "more", "than", "such", "who", "was", "were"}
    jd_words = {w for w in re.findall(r"[A-Za-z][A-Za-z0-9+#.\-]{2,}", jd.lower())
                if w not in stop and len(w) > 3}
    if not jd_words:
        return {"score": 0, "max": 0, "notes": "job description had no usable keywords"}
    resume_words = set(re.findall(r"[A-Za-z][A-Za-z0-9+#.\-]{2,}", text.lower()))
    matches = jd_words & resume_words
    missing = sorted(jd_words - resume_words)[:20]
    return {"score": round((len(matches) / len(jd_words)) * 14), "max": 14,
            "notes": f"{len(matches)}/{len(jd_words)} keywords matched",
            "missing": missing}


SCORERS: dict[str, Callable] = {
    "Contact info":       score_contact,
    "Section headers":    score_sections,
    "Quantified results": score_quantified,
    "Action verbs":       score_action_verbs,
    "Resume length":      score_length,
    "Formatting":         score_formatting,
}


def compute_score(text: str, jd: str = "") -> dict:
    """Run every scorer. Returns {ok, error, raw, summary}."""
    try:
        breakdown = []
        total, max_total = 0, 0
        for name, fn in SCORERS.items():
            r = fn(text)
            breakdown.append({"name": name, "score": r["score"],
                              "max": r["max"], "notes": r["notes"]})
            total += r["score"]
            max_total += r["max"]

        missing_keywords: list[str] = []
        kw = score_keywords(text, jd)
        if kw["max"] > 0:
            breakdown.append({"name": "Job-description keywords",
                              "score": kw["score"], "max": kw["max"],
                              "notes": kw["notes"]})
            total += kw["score"]
            max_total += kw["max"]
            missing_keywords = kw.get("missing", [])

        pct = round((total / max_total) * 100) if max_total else 0
        return {"ok": True, "error": None,
                "raw": {"breakdown": breakdown},
                "summary": {"overall_score": pct,
                            "raw_score": total,
                            "max_score": max_total,
                            "breakdown": breakdown,
                            "missing_keywords": missing_keywords}}
    except Exception as e:
        return {"ok": False, "error": str(e), "raw": {}, "summary": {}}